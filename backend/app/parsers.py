import io


def extract_text(filename: str, payload: bytes) -> str:
    if len(payload) > 5 * 1024 * 1024:
        raise ValueError("单文件不超过 5MB")
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix == "txt":
        return payload.decode("utf-8", errors="ignore")
    if suffix == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(payload))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == "docx":
        from docx import Document

        document = Document(io.BytesIO(payload))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if suffix == "xls":
        raise ValueError("请把 Excel 另存为 xlsx 后再上传")
    if suffix == "xlsx":
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        lines = []
        for sheet in book.worksheets:
            lines.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if cells:
                    lines.append(" ".join(cells))
        return "\n".join(lines)
    if suffix == "pptx":
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(payload))
        slides = []
        for index, slide in enumerate(presentation.slides, start=1):
            texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            slides.append(f"# 第{index}页\n" + "\n".join(texts))
        return "\n\n".join(slides)
    raise ValueError("只接受 PDF、Word、Excel、TXT、PPT")
